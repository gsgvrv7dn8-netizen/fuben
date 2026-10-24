"""Read-only source audit; write only the three named research reports.

AI assistance: Codex, OpenAI, 2026-10-24. Exact client model/version is
unconfirmed and must be recorded from the client before competition delivery.
No source workbook is saved, and no model is trained.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from decimal import Decimal
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


META = {"beer", "beer_id", "tasting_category_fine"}
ALIASES = {"A_ester_fruity": "A_esters_fruity", "F_ester_fruity": "F_esters_fruity", "4vg": "X4vg"}
DEFECTS = {"DMS": "dms", "X4VG": "X4vg", "X2.3.butanedione": "diacetyl", "X.E..2.nonenal": "t2n"}


def source_file(root: Path, name: str) -> Path:
    matches = list(root.rglob(name))
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one {name}; found {len(matches)}")
    return matches[0]


def number(value: object) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def rounding_tolerance(token: str) -> float:
    """Respect the CSV's actual last printed digit, including scientific notation."""
    if not token or token.upper() == "NA":
        return 0.0
    return float(Decimal(10) ** Decimal(Decimal(token).as_tuple().exponent)) / 2


def compare_transform(x: np.ndarray, raw: np.ndarray, tolerance: np.ndarray, name: str) -> dict:
    valid = np.isfinite(x) & np.isfinite(raw)
    expected = raw.copy()
    if name in {"log10", "ln"}:
        valid &= raw > 0
        expected[:] = np.nan
        expected[valid] = (np.log10 if name == "log10" else np.log)(raw[valid])
    # The second term permits only floating-point arithmetic noise, not fit error.
    matches = valid & (np.abs(x - expected) <= tolerance + 1e-12 * np.maximum(1, np.abs(expected)))
    errors = np.abs(x[valid] - expected[valid])
    return {"name": name, "valid": valid, "matches": matches,
            "n": int(valid.sum()), "matched": int(matches.sum()),
            "max_error": float(errors.max()) if errors.size else None,
            "mae": float(errors.mean()) if errors.size else None}


def card_label(card, address: str) -> str:
    cell = card[address]
    left = cell.column == 17  # Q is the left result column; AG is the right one.
    if cell.row in {10, 12, 14, 16, 23, 25, 27, 29, 36, 38, 40, 42}:
        return str(card.cell(cell.row, 5 if left else 21).value or "unknown")
    return str(card.cell(cell.row - 1, 2 if left else 18).value or "unknown")


def target_dictionary(card_path: Path, sensory: pd.DataFrame, names: pd.DataFrame) -> pd.DataFrame:
    panel = names.loc[names["dataset"] == "Trained panel"].set_index("R_name")
    if not panel.index.is_unique:
        raise ValueError("Duplicate panel R_name entries")
    workbook = openpyxl.load_workbook(card_path, data_only=False)
    card, summary = workbook["beer_1"], workbook["summary"]
    records = []
    try:
        for header in summary[1][1:]:
            original = str(header.value)
            key = ALIASES.get(original, original)
            formula = str(summary.cell(2, header.column).value)
            found = re.fullmatch(r"=beer_1!\$([A-Z]+)\$(\d+)", formula)
            if not found:
                raise ValueError(f"Unexpected summary formula: {formula}")
            reference = "".join(found.groups())
            intended = "AG90" if key == "barnyard" else reference
            entry = panel.loc[key] if key in panel.index else None
            is_target = key in sensory.columns and key not in META
            formula_errors = []
            for n in range(1, 13):
                current = str(summary.cell(n + 1, header.column).value)
                desired = f"=beer_{n}!${card[intended].column_letter}${card[intended].row}"
                if current != desired:
                    formula_errors.append(f"summary!{summary.cell(n + 1, header.column).coordinate}:{current}")
            if key.startswith(("A_", "F_")):
                structure = "Aroma" if key.startswith("A_") else "Flavour"
                structure += "/" + key.split("_")[1] + ("/overall" if key.endswith("_all") else "/descriptor_gated_by_overall")
            elif key in {"dms", "metallic", "X4vg", "t2n", "diacetyl", "stale_hops"}:
                structure = "Faults"
            elif key in {"orange", "coriander", "clove"}:
                structure = "Spiced beers"
            elif key in {"lactic", "acetic", "barnyard"}:
                structure = "Lambic/Geuze"
            else:
                structure = "General/" + key
            notes = []
            if key != original:
                notes.append("Explicit spelling/case alias; source names preserved")
            if key == "barnyard":
                notes.append("All 12 summary formulas point to Stale hops AG76; semantic card position is Barnyard AG90")
            records.append({
                "target_key": key, "record_kind": "labeled_target" if is_target else "free_text",
                "csv4_field": key if is_target else "unknown", "summary_header": original,
                "summary_header_cell": header.coordinate, "summary_reference_actual": reference,
                "card_semantic_score_cell": intended, "card_label": card_label(card, intended) if key != "comments" else "Comments:",
                "card_hierarchy": structure, "card_score_formula": str(card[intended].value or ""),
                "summary_formula_issues": json_text(formula_errors),
                "metadata_R_name": key if entry is not None else "unknown",
                "metadata_category_paper": entry["Category_Paper"] if entry is not None else "unknown",
                "metadata_category": entry["Category"] if entry is not None else "unknown",
                "metadata_datatype": entry["DataType"] if entry is not None else "unknown",
                "metadata_label": entry["CleanName_PanelFeedback"] if entry is not None else "unknown",
                "label_scale": "standardized_or_centered; inverse_parameters_unknown" if is_target else "not_numeric_target",
                "official_output_scope": "unresolved", "notes": "; ".join(notes),
            })
        for address, key, kind in [("Q45", "card_aroma_acetaldehyde", "unlabeled_sensory_item"),
                                    ("AG45", "card_flavour_acetaldehyde", "unlabeled_sensory_item"),
                                    ("AG62", "card_unlabeled_AG62", "unlabeled_formula"),
                                    ("AG94", "card_unlabeled_AG94", "unlabeled_formula")]:
            records.append({"target_key": key, "record_kind": kind, "csv4_field": "unknown",
                            "summary_header": "unknown", "summary_reference_actual": "unknown",
                            "card_semantic_score_cell": address, "card_label": card_label(card, address),
                            "card_hierarchy": "Aroma" if address == "Q45" else "Flavour" if address == "AG45" else "unknown",
                            "card_score_formula": card[address].value, "metadata_R_name": "unknown",
                            "label_scale": "no_observed_sensory_labels", "official_output_scope": "unresolved",
                            "notes": "No corresponding 4.csv trained-panel target; chemical acetaldehyde is not a sensory label" if "acetaldehyde" in key else "Formula alone does not establish a sensory target"})
        result = pd.DataFrame(records).fillna("unknown")
        mapped = set(result.loc[result.record_kind == "labeled_target", "csv4_field"])
        if mapped != set(sensory.columns) - META:
            raise ValueError("Incomplete or extra target mapping")
        if mapped != set(panel.index):
            raise ValueError("Sensory metadata target set differs from 4.csv")
        return result
    finally:
        workbook.close()


def chemical_dictionary(chem: pd.DataFrame, text: pd.DataFrame, matrix: pd.DataFrame,
                        metadata: pd.DataFrame, calibrated: pd.DataFrame) -> tuple[pd.DataFrame, float]:
    cols = [c for c in chem.columns if c not in META]
    if len(cols) != len(matrix) or matrix.shape[0] != matrix.shape[1]:
        raise ValueError("Chemical correlation matrix dimensions differ")
    if list(matrix.index) != list(matrix.columns):
        raise ValueError("Correlation row and column names differ")
    observed = chem[cols].corr(method="spearman").to_numpy()
    matrix_error = float(np.max(np.abs(observed - matrix.to_numpy(float))))
    if not np.isfinite(matrix_error) or matrix_error > 1e-12:
        raise ValueError("Cannot establish supplied matrix column order by correlation equality")
    if not metadata.field_name_R.is_unique:
        raise ValueError("Duplicate chemical metadata keys")
    meta = metadata.set_index("field_name_R")
    if not chem.beer.is_unique or not calibrated.beer.is_unique:
        raise ValueError("Nonunique sample beer key")
    if set(chem.beer) != set(calibrated.beer):
        raise ValueError("Calibrated and main sample sets differ")
    raw = calibrated.set_index("beer").reindex(chem.beer)
    records = []
    for pos, col in enumerate(cols):
        entry = meta.loc[col] if col in meta.index else None
        unit = str(entry["unit"]) if entry is not None and entry["unit"] else "unknown"
        mean = number(entry["threshold_mean"]) if entry is not None else None
        median = number(entry["threshold_median"]) if entry is not None else None
        record = {
            "feature": col, "correlation_name": matrix.columns[pos], "correlation_mapping_evidence": "whole_matrix_spearman_equality_in_supplied_order",
            "metadata_field_name_R": col if entry is not None else "unknown",
            "metadata_field_name_python": entry["field_name_python"] if entry is not None else "unknown",
            "chemical_name_source": entry["chemical_name"] if entry is not None else "unknown",
            "unit_source": unit, "unit_source_location": "8.xlsx/Sensory thresholds/unit",
            "calibration_flag_source": entry["calibrated"] if entry is not None else "unknown",
            "origin_source": entry["origin_paper"] if entry is not None else "unknown",
            "off_flavors_source": entry["off_flavors"] if entry is not None else "unknown",
            "calibrated_column_available": col in raw, "confirmed_transform": "unknown",
            "transform_scope": "unknown", "best_candidate": "unknown", "candidate_matched": 0, "candidate_n": 0,
            "candidate_max_abs_error": "unknown", "candidate_mean_abs_error": "unknown",
            "candidate_mismatch_beer": "[]", "candidate_mismatch_examples": "[]",
            "comparison_details": "{}", "raw_nonpositive_count": "unknown", "raw_missing_count": "unknown",
            "csv1_missing_count": int(chem[col].isna().sum()), "raw_nonpositive_examples": "[]",
            "threshold_mean_source": mean if mean is not None else "unknown",
            "threshold_median_source": median if median is not None else "unknown",
            "threshold_choice": "not_selected", "numeric_threshold_available": mean is not None or median is not None,
            "sensory_defect_counterpart": DEFECTS.get(col, "none"),
            "threshold_comparison_status": "unknown", "notes": "",
        }
        if col in raw:
            x = chem[col].to_numpy(float)
            r = pd.to_numeric(raw[col], errors="coerce").to_numpy(float)
            tolerance = np.array([rounding_tolerance(v) for v in text[col].fillna("")])
            options = [compare_transform(x, r, tolerance, n) for n in ("identity", "log10", "ln")]
            best = max(options, key=lambda v: v["matched"] / v["n"] if v["n"] else -1)
            full = [v for v in options if v["n"] >= 3 and v["matched"] == v["n"]]
            confirmed = full[0]["name"] if len(full) == 1 else "unknown"
            failures = np.flatnonzero(best["valid"] & ~best["matches"])
            nonpositive = np.flatnonzero(np.isfinite(r) & (r <= 0))
            record.update({
                "confirmed_transform": confirmed,
                "transform_scope": "all_finite_positive_calibrated_pairs_only" if confirmed in {"log10", "ln"} else "all_finite_pairs" if confirmed == "identity" else "unknown",
                "best_candidate": best["name"], "candidate_matched": best["matched"], "candidate_n": best["n"],
                "candidate_max_abs_error": best["max_error"], "candidate_mean_abs_error": best["mae"],
                "candidate_mismatch_beer": json_text([int(chem.beer.iloc[i]) for i in failures]),
                "candidate_mismatch_examples": json_text([{"beer": int(chem.beer.iloc[i]), "csv1": float(x[i]), "calibrated": float(r[i])} for i in failures[:8]]),
                "comparison_details": json_text({v["name"]: {k: v[k] for k in ("n", "matched", "mae", "max_error")} for v in options}),
                "raw_nonpositive_count": len(nonpositive), "raw_missing_count": int((~np.isfinite(r)).sum()),
                "raw_nonpositive_examples": json_text([{"beer": int(chem.beer.iloc[i]), "csv1": float(x[i]) if np.isfinite(x[i]) else None, "calibrated": float(r[i])} for i in nonpositive[:8]]),
            })
            if confirmed == "unknown":
                record["notes"] = "Majority agreement is not a global transformation proof; preserve mismatches"
            elif nonpositive.size:
                record["notes"] = "Positive-pair relation does not explain nonpositive-value replacement or detection limits"
        if record["numeric_threshold_available"] and unit != "unknown":
            if record["confirmed_transform"] in {"identity", "log10", "ln"}:
                record["threshold_comparison_status"] = "conditional_on_verified_transform_scope_and_threshold_choice"
            else:
                record["threshold_comparison_status"] = "threshold_present_but_csv1_scale_unresolved"
        else:
            record["threshold_comparison_status"] = "missing_threshold_or_unit"
        records.append(record)
    return pd.DataFrame(records), matrix_error


def write_findings(path: Path, targets: pd.DataFrame, chemicals: pd.DataFrame,
                   files: dict[str, Path], matrix_error: float) -> None:
    counts = chemicals.confirmed_transform.value_counts().to_dict()
    partial = chemicals.loc[(chemicals.confirmed_transform == "unknown") & chemicals.calibrated_column_available]
    lines = ["# 目标字典与化学单位/变换核验", "", "核查日期：2026-10-24。只读原始附件；未训练模型。",
             "", "## 主要结论", "",
             "- 4.csv 的 50 个感官目标均与 8.xlsx 的 Trained panel 元数据逐项对应；3.xlsx summary 通过三个显式拼写/大小写别名可覆盖这 50 项。",
             "- 评分卡还含 Aroma/Flavour 的两个 Acetaldehyde 项（Q45、AG45），4.csv 没有对应感官标签。化学 acetaldehyde 浓度不能代替它们。",
             "- summary 的 barnyard 在全部 12 张卡都错误引用 AG76（Stale hops），语义正确位置是 AG90（Barnyard/horsey）。本次只记录，没有修改原表。",
             "- AG62 和 AG94 有公式，但前者无感官项目名称，后者邻接 Comments。公式存在并不构成额外有标签目标。",
             "- 原始感官 0–6 尺度和官方输出范围仍未解决；此字典用于对齐，不是自行宣布官方要求。",
             "", "## 目标的层次与计数", "",
             "Aroma 与 Flavour 的 Malt/Hops/Esters 各有一个总强度及四个子描述词，共 2×3×5=30 个 A_/F_ 标签。24 个子描述词公式不是独立选 0–6，而是 IF(未勾选,0,对应总强度)，元数据也标为 categorical；这是原评分卡编码结构，不能把标准化后的值直接认作分类等级。",
             "另有 7 个一般感官项（body、alcohol、bitternes、sweetness、co2、acidity、aftertaste）、6 个缺陷项、6 个风格附加项和 overall，共 20 项，合计 50。卡面额外的两个 Acetaldehyde 项使有名称的感官评分位置达到 52。元数据按 Category_Paper 归为 Aroma 15、Taste 30、Palate 4、Overall 1，与卡面左右排版不是同一分类。",
             "题面“约 40”没有与这几个集合精确对应。30 个 A_/F_ 加 7 个一般项为 37，再加未入标签的两个 Acetaldehyde 为 39，只是可能解释近似数量的算术组合，并非官方意图证据。必须查官方说明后再决定输出范围，不能自行删减至 40。",
             "", "## 化学映射与变换覆盖", "",
             f"- 231 个 1.csv 特征全部通过 field_name_R 找到阈值元数据；2.csv 的231个显示名称按现有顺序重算 Spearman 的全矩阵最大误差为 {matrix_error:.17g}。这证实当前附件顺序对应，但字典仍保存两套名字。",
             f"- {int(chemicals.calibrated_column_available.sum())}/231 个特征在校准化学表中有同名列；{int((chemicals.unit_source != 'unknown').sum())}/231 个特征有非空来源单位。",
             f"- 严格逐值核验：{counts.get('log10', 0)} 列为 log10（仅正值有效配对），{counts.get('identity', 0)} 列为 identity，{counts.get('ln', 0)} 列为 ln，{counts.get('unknown', 0)} 列仍为 unknown。",
             f"- unknown 中 {len(partial)} 列虽有校准对照、且多数点接近 identity，但存在实质差异；其余列缺少校准对照，不能仅凭正负号或范围判断。CSV 保留每列失败 beer 编号和差异样例。",
             f"- {int(chemicals.numeric_threshold_available.sum())}/231 个特征在来源表至少有一个数值阈值；有阈值不等于已可无条件做缺陷约束。",
             "",
             "比较容差取 1.csv 原始文本末位精度的半个单位（包括科学计数法）加浮点运算噪声 1e-12×max(1,|对照|)，避免把 8.95e-05 与 8.94815e-05 这种显示舍入误判为数据变换。只有一个候选变换在其适用域的所有有效配对都通过、且至少三个配对时才标 confirmed。这里的确证限于该有限数据及有效域，不证明未观测数据上的处理机制。",
             "",
             "37 个 log10 列包含 pH、ethanol、bitterness、color、CO2 等。非正校准值对应 1.csv 有限数值，说明存在替换/预处理；不能把这些值逆变换后当真实实测浓度，也不能把它们自动叫检测下限。校准表本身还存在负值及缺失，因此也不是无需核验的绝对真值。",
             "", "## 五项重点缺陷", "",
             "| 项目 | 化学列 / 感官列 | 来源单位与阈值 | 当前可用性 |", "|---|---|---|---|",
             "| DMS | DMS / dms | ug/L；均值17.25，中位19 | 249个正值配对为log10；beer=10校准值为0但输入有限，须处理替换语义；阈值均值/中位尚未选择 |",
             "| 4VG | X4VG / X4vg | mg/L；均值0.206333333，中位0.2 | 82个正值配对为log10；168个校准零值对应固定替换值，不能当168个实测正浓度 |",
             "| Diacetyl | X2.3.butanedione / diacetyl | mg/L；均值0.29925，中位0.09 | 238/241个有效配对符合identity，但beer 151、156、34明显不同；全列变换仍unknown，需处理差异及9个输入缺失 |",
             "| trans-2-nonenal | X.E..2.nonenal / t2n | mg/L；两种阈值均缺失 | 247/249个配对符合identity，beer158、130不同；阈值和全列尺度未闭合 |",
             "| Stale hops | 未找到单独同名化学测量列 / stale_hops | 缺少单一物质与单位阈值 | 是现有感官目标；aroma_hops_sum为总酒花香气汇总，不能自动充当其化学代理 |",
             "",
             "阈值 mean 和 median 是来源数据的两个汇总版本，不是可任选后隐瞒选择的同一常数；尤其 Diacetyl 两值差异较大。目标缺陷感官分数与相应化学浓度是不同变量，不能用标准化 dms/diacetyl 分数直接比较 ug/L 或 mg/L 阈值。",
             "",
             "## 尚不能识别的部分", "",
             "- 未从本地附件取得明确检测下限/定量下限或零值替换规范；来源字段中 NA/空值保留 unknown，不填文献常识值。",
             "- 无校准对照列的单位即使给出，也不足以证明 1.csv 的存储尺度；单位与变换分别记录。",
             "- 8.xlsx 部分显示名存在字符编码问题，以及字段名与显示化学名不完全相符的条目；保留原文，不能用模糊名称匹配悄悄改名。",
             "- 32列多数identity而少数不一致可能源于处理差异、数据修订或其他原因；本核查不猜测原因，也不自动以校准表覆盖比赛输入。",
             "", "## 复现与文件", "",
             "运行：`python research/audit_units_targets.py --source D:\\MathorCup`。需要 pandas、numpy、openpyxl。脚本只写同目录的 target-dictionary.csv、chemical-units.csv、target-unit-findings.md；原始附件保持只读。",
             "", "- target-dictionary.csv：50个有标签目标、2个卡片额外感官项、2个无标签公式位置、1个自由文本位置，共55行；官方输出范围均标 unresolved。",
             "- chemical-units.csv：231行，记录来源字段/单位、每种变换对比、失败样本、阈值及条件可用性。",
             "", "来源 SHA-256：", ""]
    for name, source in files.items():
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        lines.append(f"- `{name}`：`{digest}`；`{source}`")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=Path(r"D:\MathorCup"))
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()
    files = {name: source_file(args.source, name) for name in ("1.csv", "2.csv", "3.xlsx", "4.csv", "8.xlsx")}
    chem = pd.read_csv(files["1.csv"])
    lexical = pd.read_csv(files["1.csv"], dtype=str, keep_default_na=False)
    matrix = pd.read_csv(files["2.csv"], index_col=0)
    sensory = pd.read_csv(files["4.csv"])
    metadata = pd.read_excel(files["8.xlsx"], sheet_name="Sensory thresholds", keep_default_na=False)
    names = pd.read_excel(files["8.xlsx"], sheet_name="Sensory names", keep_default_na=False)
    calibrated = pd.read_excel(files["8.xlsx"], sheet_name="Calibrated chemical dataset")
    targets = target_dictionary(files["3.xlsx"], sensory, names)
    chemicals, error = chemical_dictionary(chem, lexical, matrix, metadata, calibrated)
    args.output.mkdir(parents=True, exist_ok=True)
    targets.to_csv(args.output / "target-dictionary.csv", index=False, encoding="utf-8-sig")
    chemicals.to_csv(args.output / "chemical-units.csv", index=False, encoding="utf-8-sig")
    write_findings(args.output / "target-unit-findings.md", targets, chemicals, files, error)
    print(json_text({"target_rows": len(targets), "chemical_rows": len(chemicals),
                     "transforms": chemicals.confirmed_transform.value_counts().to_dict()}))


if __name__ == "__main__":
    main()
